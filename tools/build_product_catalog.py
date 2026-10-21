#!/usr/bin/env python3
"""Build product-catalog.xlsx from a saved instax.in scrape.

The site is a Shopify store, so every product page carries the full product
and variant data in its embedded ``initData`` JSON. This tool reads those
product entries, groups variants by price, and writes the Excel list the app
loads for extraction guidance and price verification.

Usage:
    python tools/build_product_catalog.py --source "C:/path/to/scraped_site/pages"
    python tools/build_product_catalog.py --source ... --output product-catalog.xlsx

Requires: openpyxl (pip install openpyxl)
"""
import argparse
import glob
import json
import os
import sys
from collections import Counter, defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["Product Name", "Price (INR)", "SKU", "Category", "Variants", "Product URL"]
SKIP_VARIANT_TITLES = {"", "default title"}
PRODUCT_URL_BASE = "https://www.instax.in"


def extract_json_object(text, start):
    """Return the balanced {...} object starting at *start* (string-aware)."""
    depth = 0
    in_string = False
    escaped = False
    for index in range(start, len(text)):
        char = text[index]
        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == '"':
                in_string = False
        else:
            if char == '"':
                in_string = True
            elif char == "{":
                depth += 1
            elif char == "}":
                depth -= 1
                if depth == 0:
                    return text[start:index + 1]
    return None


def page_products(path):
    """Extract the ``initData.products`` list from one scraped product page."""
    text = open(path, encoding="utf-8", errors="replace").read()
    marker = text.find("initData:")
    if marker < 0:
        return []
    start = text.find('{"shop"', marker)
    if start < 0:
        return []
    raw = extract_json_object(text, start)
    if not raw:
        return []
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        return []
    return data.get("products") or []


def collect_products(pages_dir):
    """Merge every product page into handle -> {title, type, url, variants}."""
    products = {}
    warnings = []
    files = sorted(glob.glob(os.path.join(pages_dir, "products__*.html")))
    if not files:
        raise SystemExit(f"No products__*.html pages found in {pages_dir}")

    for path in files:
        for raw in page_products(path):
            url = str(raw.get("url") or "").strip()
            handle = url.strip("/").split("/")[-1].lower()
            if not handle:
                continue
            entry = products.setdefault(handle, {
                "titles": Counter(),
                "types": Counter(),
                "url": PRODUCT_URL_BASE + url if url.startswith("/") else url,
                "variants": {},
            })
            title = " ".join(str(raw.get("title") or "").split())
            if title:
                entry["titles"][title] += 1
            product_type = " ".join(str(raw.get("type") or "").split())
            if product_type:
                entry["types"][product_type] += 1
            for variant in raw.get("variants") or []:
                price = (variant.get("price") or {}).get("amount")
                try:
                    price = round(float(price), 2)
                except (TypeError, ValueError):
                    continue
                if price <= 0:
                    continue
                variant_title = " ".join(str(variant.get("title") or "").split())
                sku = " ".join(str(variant.get("sku") or "").split())
                key = (variant_title, price, sku)
                entry["variants"][key] = entry["variants"].get(key, 0) + 1

    for handle, entry in products.items():
        titles = entry["titles"]
        if len(titles) > 1:
            warnings.append(f"{handle}: multiple titles {dict(titles)}; using the most common")
        entry["title"] = titles.most_common(1)[0][0] if titles else handle
        entry["type"] = entry["types"].most_common(1)[0][0] if entry["types"] else ""
        prices_by_variant = defaultdict(set)
        for (variant_title, price, _sku) in entry["variants"]:
            prices_by_variant[variant_title].add(price)
        for variant_title, prices in prices_by_variant.items():
            if len(prices) > 1:
                warnings.append(f"{handle} / {variant_title!r}: multiple prices {sorted(prices)}; keeping all")
    return products, warnings


def row_for_group(name, price, variants, skus, product):
    variant_titles = sorted({title for title in variants if title.lower() not in SKIP_VARIANT_TITLES})
    return {
        "name": name,
        "price": price,
        "sku": " / ".join(sorted({sku for sku in skus if sku})),
        "category": product["type"],
        "variants": " / ".join(variant_titles),
        "url": product["url"],
    }


def catalog_rows(products):
    rows = []
    for product in products.values():
        variants = sorted(product["variants"].items(), key=lambda item: (item[0][1], item[0][0].lower()))
        prices = sorted({price for (_, price, _sku), _count in variants})
        if not prices:
            continue
        base = product["title"]
        if len(prices) == 1:
            rows.append(row_for_group(base, prices[0], [item[0][0] for item in variants],
                                      [item[0][2] for item in variants], product))
            continue
        for price in prices:
            group = [item[0] for item in variants if item[0][1] == price]
            labels = sorted({title for (title, _price, _sku) in group if title.lower() not in SKIP_VARIANT_TITLES})
            suffix = f" - {' / '.join(labels)}" if labels else ""
            rows.append(row_for_group(base + suffix, price, [title for (title, _p, _s) in group],
                                      [sku for (_t, _p, sku) in group], product))
    rows.sort(key=lambda row: (row["name"].casefold(), row["price"]))
    unique = {}
    for row in rows:
        unique.setdefault((row["name"].casefold(), row["price"]), row)
    return sorted(unique.values(), key=lambda row: (row["name"].casefold(), row["price"]))


def style_sheet(sheet, headers, widths):
    header_fill = PatternFill("solid", fgColor="1F4E3D")
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column)
        cell.value = header
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(column)].width = widths[column - 1]
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    sheet.row_dimensions[1].height = 22


def write_workbook(rows, output):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append([row["name"], row["price"], row["sku"], row["category"], row["variants"], row["url"]])
        sheet.cell(row=sheet.max_row, column=2).number_format = "#,##0.00"
        sheet.cell(row=sheet.max_row, column=6).font = Font(color="1F5FAF", underline="single")
        sheet.cell(row=sheet.max_row, column=6).hyperlink = row["url"]
    style_sheet(sheet, HEADERS, [58, 14, 20, 18, 42, 62])
    for row_number in range(2, sheet.max_row + 1):
        sheet.cell(row=row_number, column=5).alignment = Alignment(wrap_text=True, vertical="top")

    readme = workbook.create_sheet("Read Me")
    readme.column_dimensions["A"].width = 110
    readme["A1"] = "Instax product and price list"
    readme["A1"].font = Font(size=16, bold=True)
    lines = [
        "",
        f"{len(rows)} products, prices in Indian rupees (INR). Generated from instax.in on {date.today().isoformat()}.",
        "",
        "How to update:",
        "1. Edit the Price (INR) column, or add a new row under the Products sheet.",
        "2. Keep the Product Name and Price (INR) headers unchanged — extra columns are ignored.",
        "3. Save this file, then reload the app. For the deployed site, redeploy the frontend so it serves the new file.",
        "4. The app uses these names and prices to guide extraction and verify scanned rates.",
        "",
        "Notes:",
        "• Products with variants priced differently (for example Mini 99 Black/Silver) have one row per price.",
        "• Deleting or renaming the Products sheet will stop the app from loading the list.",
        "• This file is served with the web app. If the app is public, anyone can download it — do not put confidential purchase costs here.",
        "• Rerunning tools/build_product_catalog.py overwrites manual edits.",
    ]
    for offset, line in enumerate(lines, start=2):
        readme.cell(row=offset, column=1, value=line)
    readme["A5"].font = Font(bold=True)
    readme["A11"].font = Font(bold=True)

    workbook.save(output)


def main():
    parser = argparse.ArgumentParser(description="Build product-catalog.xlsx from a scraped instax.in site.")
    parser.add_argument("--source", required=True, help="Folder containing the scraped pages/products__*.html files")
    parser.add_argument("--output", default="product-catalog.xlsx", help="Output .xlsx path (default: product-catalog.xlsx)")
    args = parser.parse_args()
    if not os.path.isdir(args.source):
        raise SystemExit(f"Source folder not found: {args.source}")

    products, warnings = collect_products(args.source)
    rows = catalog_rows(products)
    if not rows:
        raise SystemExit("No priced products found — check the source folder.")
    write_workbook(rows, args.output)
    for warning in warnings:
        print("WARN:", warning, file=sys.stderr)
    print(f"Wrote {len(rows)} rows for {len(products)} products to {args.output}")


if __name__ == "__main__":
    main()
